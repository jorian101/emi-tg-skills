#!/usr/bin/env python3
"""Verifica un entregable contra el formato oficial de un docente: hoja, márgenes, columnas, páginas y, sobre todo, el ORDEN.

Uso:
  verificar_formato.py <archivo> --tipo articulo|biptico|triptico|diapositivas|tg [--docente <código|slug>] [--vault <vault>]
                       [--catalogo <carpeta>] [--formato <ruta/formato.yaml>]
Los formatos viven en `<catalogo>/formatos/<docente>/formato.yaml` (con su plantilla vacía). Qué formato se usa:
  1. `--formato` (un formato.yaml cualquiera);
  2. `--docente`;
  3. con `--vault`, el del primer evaluador vinculado que define ese tipo (docente de TG, tutor, revisor 1, revisor 2);
  4. si solo otro docente lo define, se usa avisando; si hay varios, se listan y se pide `--docente` (no se mezclan).
Sale 1 si hay algún desvío; lo que el docente dio solo como ejemplo (`obligatorio: false`) se informa como aviso.
"""

from __future__ import annotations

import argparse
import re
import sys
import zipfile
from pathlib import Path

import yaml
from catalogo_lib import catalogo_default, frontmatter, normalizar
from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.util import Emu

TOL_CM = 0.15
ROLES = ("Docente_TG.md", "Tutor.md", "Revisor_1.md", "Revisor_2.md")
ENUM = re.compile(r"^(?:[ivxlc]+|\d+(?:\.\d+)*|[a-z])[.)]\s+", re.IGNORECASE)


def limpiar(texto: str) -> str:
    t = normalizar(re.sub(r"\s+", " ", texto)).strip(" :.-—")
    return ENUM.sub("", t)


# ------------------------------------------------------------------ elegir el formato
def formatos_del_catalogo(cat: Path, tipo: str) -> dict[str, Path]:
    out = {}
    for f in sorted((cat / "formatos").glob("*/formato.yaml")):
        if tipo in (yaml.safe_load(f.read_text(encoding="utf-8")) or {}):
            out[f.parent.name] = f
    return out


def evaluadores(vault: Path) -> list[str]:
    cods = []
    for rol in ROLES:
        p = vault / "wiki/revisores" / rol
        d = str(frontmatter(p.read_text(encoding="utf-8")).get("docente") or "").strip("[]") if p.is_file() else ""
        if d and d != "por-asignar":
            cods.append(d)
    return cods


def elegir(a, tipo: str) -> tuple[Path, str]:
    if a.formato:
        return a.formato, ""
    disponibles = formatos_del_catalogo(a.catalogo, tipo)
    if not disponibles:
        sys.exit(f"El catálogo ({a.catalogo}) no trae ningún formato de «{tipo}»: bajá el catálogo público o pasá --formato.")
    if a.docente:
        if a.docente not in disponibles:
            sys.exit(f"{a.docente} no define «{tipo}». Lo definen: {', '.join(disponibles)}")
        return disponibles[a.docente], ""
    if a.vault:
        for cod in evaluadores(a.vault):
            if cod in disponibles:
                return disponibles[cod], ""
    if len(disponibles) == 1:
        cod = next(iter(disponibles))
        return disponibles[cod], f"el formato de «{tipo}» es de {cod}, que no figura entre tus evaluadores: confirmá que es el que te corresponde"
    sys.exit(f"Hay varios formatos de «{tipo}» ({', '.join(disponibles)}) y ninguno de tus evaluadores lo define: elegí con --docente.")


# ------------------------------------------------------------------ comprobaciones sueltas
class Informe:
    def __init__(self) -> None:
        self.fallos: list[str] = []
        self.avisos: list[str] = []
        self.ok: list[str] = []

    aviso_cm = False  # formato que es solo un ejemplo: las medidas fuera de rango son avisos

    def cm(self, nombre: str, real: float, esperado: float) -> None:
        (self.ok if abs(real - esperado) <= TOL_CM else self.avisos if self.aviso_cm else self.fallos).append(
            f"{nombre}: {real:.2f} cm" + ("" if abs(real - esperado) <= TOL_CM else f" (el formato pide {esperado:g})"))

    def check(self, cond: bool, bien: str, mal: str, aviso: bool = False) -> None:
        (self.ok if cond else (self.avisos if aviso else self.fallos)).append(bien if cond else mal)


def margenes(inf: Informe, seccion, esperados: dict) -> None:
    for clave, attr in (("sup", "top_margin"), ("inf", "bottom_margin"), ("izq", "left_margin"), ("der", "right_margin")):
        if clave in esperados:
            inf.cm(f"margen {clave}", getattr(seccion, attr).cm, float(esperados[clave]))


def paginas_word(path: Path) -> int | None:
    with zipfile.ZipFile(path) as z:
        m = re.search(r"<Pages>(\d+)</Pages>", z.read("docProps/app.xml").decode("utf-8", "ignore")) if "docProps/app.xml" in z.namelist() else None
    return int(m[1]) if m else None


def textos_docx(doc) -> list[str]:
    """Párrafos en orden de documento (incluye cuadros de texto), sin repetidos consecutivos (los cuadros duplican su respaldo)."""
    out: list[str] = []
    for p in doc.element.body.iter(qn("w:p")):
        t = limpiar("".join(x.text or "" for x in p.iter(qn("w:t"))))
        if t and (not out or out[-1] != t):
            out.append(t)
    return out


def comprobar_orden(inf: Informe, textos: list[str], orden: list[dict], etiqueta: str, aviso: bool = False) -> None:
    pos, faltan, fuera = -1, [], []
    for item in orden:
        if item.get("reemplazable"):  # marcador que el estudiante sustituye por su dato (p. ej. el título)
            continue
        titulos = [limpiar(t) for t in item["titulos"]]
        enc = next((i for i in range(pos + 1, len(textos)) if any(textos[i].startswith(t) for t in titulos)), None)
        if enc is not None:
            pos = enc
            continue
        antes = any(any(textos[i].startswith(t) for t in titulos) for i in range(pos + 1))
        (fuera if antes else faltan).append(item["titulos"][0])
    if not faltan and not fuera:
        inf.ok.append(f"{etiqueta}: las {len(orden)} partes están y en el orden del formato")
    else:
        if faltan:
            inf.check(False, "", f"{etiqueta}: faltan {', '.join(faltan)}", aviso)
        if fuera:
            inf.check(False, "", f"{etiqueta}: fuera de orden {', '.join(fuera)}", aviso)


# ------------------------------------------------------------------ tipos
def verificar_docx(path: Path, spec: dict, inf: Informe, aviso: bool) -> None:
    doc = Document(str(path))
    s = doc.sections[0]
    ancho, alto = (s.page_width.cm, s.page_height.cm)
    apaisada = spec.get("orientacion") == "apaisada"
    if apaisada:
        ancho, alto = max(ancho, alto), min(ancho, alto)
    if spec.get("hoja") == "carta":
        inf.cm("ancho de hoja", ancho, 27.94 if apaisada else 21.59)
        inf.cm("alto de hoja", alto, 21.59 if apaisada else 27.94)
    margenes(inf, s, spec.get("margenes_cm", {}))
    if spec.get("columnas"):
        cols = [int(c.get(qn("w:num")) or 1) for sec in doc.sections for c in sec._sectPr.xpath("./w:cols")]
        inf.check(spec["columnas"] in cols, f"columnas: el cuerpo va a {spec['columnas']}", f"columnas: ninguna sección va a {spec['columnas']} columnas", aviso)
    paginas = paginas_word(path)
    for clave, texto in (("max_hojas", "hojas"), ("paginas", "páginas")):
        if spec.get(clave) and paginas:
            ok = paginas <= spec[clave] if clave == "max_hojas" else paginas == spec[clave]
            inf.check(ok, f"{texto}: {paginas}", f"{texto}: {paginas} (el formato pide {'como máximo ' if clave == 'max_hojas' else ''}{spec[clave]})", aviso)
    if spec.get("tablas"):
        reales = [[len(t.rows), len(t.columns)] for t in doc.tables]
        inf.check(reales[: len(spec["tablas"])] == spec["tablas"], "tablas de paneles como en el formato", f"tablas de paneles {reales}; el formato usa {spec['tablas']}", aviso)
    if spec.get("orden"):
        comprobar_orden(inf, textos_docx(doc), spec["orden"], "orden", aviso)


def estilo_efectivo(p, attr: str):
    estilo = p.style
    while estilo is not None:
        v = getattr(estilo.font, attr, None) if attr in ("name", "size") else getattr(estilo.paragraph_format, attr, None)
        if v is not None:
            return v
        estilo = estilo.base_style
    return None


def verificar_tg(path: Path, spec: dict, inf: Informe) -> None:
    doc = Document(str(path))
    s = doc.sections[0]
    inf.cm("ancho de hoja", s.page_width.cm, 21.59)
    margenes(inf, s, spec.get("margenes_cm", {}))
    cuerpo = [p for p in doc.paragraphs if len(p.text) > 120 and not p.style.name.lower().startswith(("toc", "índice", "indice"))]
    if not cuerpo:
        inf.avisos.append("no encontré párrafos de cuerpo para medir letra e interlineado")
        return
    from collections import Counter
    fuente = Counter((p.runs[0].font.name if p.runs and p.runs[0].font.name else estilo_efectivo(p, "name")) for p in cuerpo).most_common(1)[0][0]
    tam = Counter((p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else (estilo_efectivo(p, "size").pt if estilo_efectivo(p, "size") else None)) for p in cuerpo).most_common(1)[0][0]
    interlineado = Counter((p.paragraph_format.line_spacing or estilo_efectivo(p, "line_spacing")) for p in cuerpo).most_common(1)[0][0]
    inf.check(fuente == spec["fuente"], f"letra del cuerpo: {fuente}", f"letra del cuerpo: {fuente} (el formato pide {spec['fuente']})")
    inf.check(tam == spec["tamano_pt"], f"tamaño del cuerpo: {tam} pt", f"tamaño del cuerpo: {tam} pt (el formato pide {spec['tamano_pt']})")
    inf.check(interlineado == spec["interlineado"], f"interlineado: {interlineado}", f"interlineado: {interlineado} (el formato pide {spec['interlineado']})")
    if spec.get("sangria_primera_linea") is False:
        con = sum(1 for p in cuerpo if (p.paragraph_format.first_line_indent or estilo_efectivo(p, "first_line_indent") or 0) > 0)
        inf.check(con == 0, "sin sangría en el cuerpo", f"{con} párrafos con sangría de primera línea (el formato pide ninguna)")


def titulo_diapositiva(s) -> str:
    for f in s.shapes:
        if f.has_text_frame and f.text_frame.text.strip() and "Marcador de contenido" not in f.name:
            t = limpiar(f.text_frame.text.split("\n")[0])
            if not t.isdigit():
                return t
    return ""


def verificar_pptx(path: Path, spec: dict, inf: Informe) -> None:
    prs = Presentation(str(path))
    ancho, alto = Emu(prs.slide_width).cm, Emu(prs.slide_height).cm
    if spec.get("tam_cm"):
        inf.cm("ancho de diapositiva", ancho, spec["tam_cm"][0])
        inf.cm("alto de diapositiva", alto, spec["tam_cm"][1])
    visibles = [s for s in prs.slides if s._element.get("show") != "0"]
    titulos = [titulo_diapositiva(s) for s in visibles]
    primera: dict[str, int] = {}
    actual = None
    asignada: list[str | None] = []
    for i, t in enumerate(titulos):
        clave = next((it["clave"] for it in spec["orden"] if any(t.startswith(limpiar(x)) for x in it["titulos"])), None)
        if clave and clave not in primera:
            primera[clave] = i
        actual = clave or actual
        asignada.append(actual if i else None)
    faltan = [it["titulos"][0] for it in spec["orden"] if it["clave"] not in primera]
    if i := len(visibles):
        inf.check(not faltan, f"orden: las {len(spec['orden'])} diapositivas del formato están ({i} visibles)", f"faltan diapositivas: {', '.join(faltan)}")
    en_orden = [primera[it["clave"]] for it in spec["orden"] if it["clave"] in primera]
    inf.check(en_orden == sorted(en_orden), "orden: van en la secuencia del formato", "orden: hay diapositivas fuera de la secuencia del formato")
    inf.check("gracias" not in primera or primera["gracias"] == len(visibles) - 1, "«Gracias» cierra la presentación", "«Gracias» no es la última diapositiva visible")
    if spec.get("max_marco_practico") and "marco_practico" in primera:
        n = sum(1 for a in asignada if a == "marco_practico")
        inf.check(n <= spec["max_marco_practico"], f"marco práctico: {n} diapositivas", f"marco práctico: {n} diapositivas (el formato pide como máximo {spec['max_marco_practico']})")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("archivo", type=Path)
    ap.add_argument("--tipo", required=True, choices=["articulo", "biptico", "triptico", "diapositivas", "tg"])
    ap.add_argument("--docente")
    ap.add_argument("--vault", type=Path)
    ap.add_argument("--catalogo", type=Path, default=catalogo_default())
    ap.add_argument("--formato", type=Path)
    a = ap.parse_args()
    tipo_clave = "tg_word" if a.tipo == "tg" else a.tipo
    ruta, nota = elegir(a, tipo_clave)
    datos = yaml.safe_load(ruta.read_text(encoding="utf-8"))
    spec = datos[tipo_clave]
    aviso = spec.get("obligatorio") is False
    inf = Informe()
    inf.aviso_cm = aviso
    if nota:
        inf.avisos.append(nota)
    if aviso:
        inf.avisos.append("este formato es un ejemplo del docente, no una norma: los desvíos son avisos")
    if a.tipo == "diapositivas":
        verificar_pptx(a.archivo, spec, inf)
    elif a.tipo == "tg":
        verificar_tg(a.archivo, spec, inf)
    else:
        verificar_docx(a.archivo, spec, inf, aviso)
    print(f"Formato: {ruta.parent.name}/{ruta.name} · {a.tipo} · {a.archivo.name}")
    for linea in inf.ok:
        print(f"  ok     {linea}")
    for linea in inf.avisos:
        print(f"  aviso  {linea}")
    for linea in inf.fallos:
        print(f"  FALLA  {linea}")
    return 1 if inf.fallos else 0


if __name__ == "__main__":
    sys.exit(main())
