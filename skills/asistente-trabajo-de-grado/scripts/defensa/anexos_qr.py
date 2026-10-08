#!/usr/bin/env python3
"""Genera el QR de cada anexo (a partir de su link de Drive) y arma QR-ANEXOS.docx (solo QRs).

Lee el manifiesto (ANEXOS-QR/manifiesto-anexos.json), crea un PNG de QR por anexo en
ANEXOS-QR/qr/, y escribe QR-ANEXOS.docx con, por anexo: caratula «ANEXO X: TITULO» + su QR + el link.

Uso: uv run --with qrcode --with pillow --with python-docx python scripts/anexos_qr.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import qrcode
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import config  # noqa: E402

_CFG = config.cargar()
BASE = config.clave(_CFG, "anexos.dir", "Es la carpeta de los anexos con QR (manifiesto, documento y salidas).")
DOCUMENTO = BASE / _CFG["anexos"].get("documento", "anexos-qr.docx")
MANIFIESTO = BASE / "manifiesto-anexos.json"
QR_DIR = BASE / "qr"
SALIDA = BASE / "QR-ANEXOS.docx"


def main() -> None:
    datos = json.loads(MANIFIESTO.read_text(encoding="utf-8"))
    anexos = datos["anexos"]
    QR_DIR.mkdir(parents=True, exist_ok=True)

    doc = Document()
    sec = doc.sections[0]
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Cm(2))

    for i, a in enumerate(anexos):
        link = a.get("drive_link") or ""
        letra = a["letra"]
        titulo = a["titulo"]

        # 1) imagen QR
        img = qrcode.make(link)
        ruta_qr = QR_DIR / f"ANEXO-{letra}.png"
        img.save(ruta_qr)

        # 2) entrada en el docx
        car = doc.add_paragraph()
        car.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = car.add_run(f"ANEXO {letra}: {titulo.upper()}")
        r.bold = True
        r.font.size = Pt(13)

        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(ruta_qr), width=Cm(5))

        pie = doc.add_paragraph()
        pie.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = pie.add_run(link)
        rr.font.size = Pt(8)
        rr.italic = True

        if i < len(anexos) - 1:
            doc.add_page_break()

    doc.save(SALIDA)
    print(f"ok {SALIDA} ({len(anexos)} QRs)")


if __name__ == "__main__":
    main()
