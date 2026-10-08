#!/usr/bin/env python3
"""Manual de usuario en Word con el formato del TG, desde un Markdown con marcas y las capturas de la interfaz.

Uso: uv run --with python-docx --with pillow --with pyyaml python manual.py manual.yaml
Modelo comentado: assets/defensa/manual.example.yaml; guía: references/defensa/manual.md.

- Base: `plantilla` (el Word del manual que el autor ajustó a mano). Se conservan su portada, índices, carátulas
  y estilos; el contenido que sigue a cada H1 cuyo texto coincide con un capítulo del .md se reemplaza. Sin
  plantilla, sale un documento nuevo con un H1 por capítulo y los estilos de Word.
- Texto: `fuente` (.md). `## CAPÍTULO N: Título` abre un capítulo; `###`, `####`, `#####` son H2 (mayúsculas),
  H3 y H4 (tipo título); `1. ` pasos numerados que reinician en cada lista; `- ` viñetas.
  Marcas: {{FIG: captura | título | nota}}, {{TABLA-NOMBRE}} (de `tablas:`), {{TABLA-GLOSARIO}} (de `glosario:`),
  {{CONSEJO: …}}, {{ATENCION: …}}, {{CODIGO: …}}. En el texto, {{REF: captura}} y {{REF-TABLA: NOMBRE}} pasan a
  «Figura N» / «Tabla N» por orden de aparición: reordenar secciones no rompe las citas.
- Figuras: la captura con sus marcadores rojos numerados (`marcadores:` o `<capturas>/marcadores.json`, que
  escribe capturas.py) y su recorte (`recortes:`). Toda tabla y figura lleva título arriba y «Nota. … Elaboración
  propia, AÑO.» abajo, como el TG.
- Falla (sin escribir nada) si falta una captura, si una referencia no existe o si hay dos títulos seguidos sin
  párrafo entre ellos.
"""

from __future__ import annotations

import itertools
import json
import re
import sys
import unicodedata
from pathlib import Path

import yaml
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from PIL import Image, ImageDraw, ImageFont

ESTILOS = {  # rol -> estilo del Word del autor; si la plantilla no lo tiene, se usa el de Word (segundo valor)
    "h1": ("H1", "Heading 1"), "h2": ("H2", "Heading 2"), "h3": ("H3", "Heading 3"), "h4": ("H4", "Heading 4"),
    "cuerpo": ("Normal oficial", "Normal"), "nota": ("Nota de fuente", "Normal"),
    "tabla_titulo": ("Tabla", "Caption"), "figura_titulo": ("Figura", "Caption"),
    "pasos": ("Numeracion", "List Number"), "vinieta": ("Viñeta sin negrilla", "List Bullet"),
    "codigo": ("SIMPLE cubo", "Normal"), "glosario_termino": ("Viñeta con negrilla", "List Bullet"),
    "glosario_def": ("Normal con sangría", "Normal"), "tabla": ("Tabladelista1clara1", "Table Grid"),
}
MENORES = {"a", "al", "con", "de", "del", "el", "en", "la", "las", "lo", "los", "o", "para", "por",
           "sin", "su", "sus", "un", "una", "y"}


class Manual:
    def __init__(self, cfg: dict, base: Path):
        r = lambda k, d=None: (base / cfg[k]).resolve() if cfg.get(k) else d
        self.capturas, self.figuras = r("capturas", base / "capturas"), r("figuras", base / "figuras")
        self.fuente, self.salida, self.plantilla = r("fuente"), r("salida", base / "manual-de-usuario.docx"), r("plantilla")
        self.anio = cfg.get("anio", "AÑO")
        self.ancho = cfg.get("ancho_figura_cm", 15)
        self.tablas, self.glosario = cfg.get("tablas", {}), cfg.get("glosario", [])
        self.recortes = cfg.get("recortes", {})
        self.marcadores = {k: [tuple(p) for p in v] for k, v in cfg.get("marcadores", {}).items()}
        medidos = self.capturas / "marcadores.json"
        if medidos.is_file():  # los mide capturas.py al capturar: se leen directo, sin copiar coordenadas
            self.marcadores.update({k: [tuple(p) for p in v] for k, v in
                                    json.loads(medidos.read_text(encoding="utf-8")).items() if v})
        self.nombres = {**{k: v[0] for k, v in ESTILOS.items()}, **cfg.get("estilos", {})}
        self.doc = Document(str(self.plantilla)) if self.plantilla else Document()

    # ---------- estilos y párrafos ----------
    def estilo(self, rol: str):
        nombre = self.nombres[rol]
        for s in self.doc.styles:
            if nombre in (s.name, s.style_id):
                return s
        return self.doc.styles[ESTILOS[rol][1]]

    def parrafo(self, rol="cuerpo", centro=False):
        p = self.doc.add_paragraph(style=self.estilo(rol))
        if centro:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return p

    def nota(self, texto: str) -> None:
        n = self.parrafo("nota")
        n.add_run("Nota.").italic = True
        n.add_run(f" {texto} Elaboración propia, {self.anio}.")

    @staticmethod
    def titulo_tg(texto: str) -> str:
        """Tipo título como el TG (artículos y preposiciones en minúscula), sin tocar lo que va entre paréntesis."""
        base, par, resto = texto.partition(" (")
        pal = base.split()
        base = " ".join(p if (i and p.lower() in MENORES) or not p[:1].isalpha() or p.isupper()
                        else p[:1].upper() + p[1:] for i, p in enumerate(pal))
        return base + (f" ({resto}" if par else "")

    @staticmethod
    def enriquecido(p, texto: str) -> None:
        """**negrita** y `código` dentro de un párrafo; tras una etiqueta «**X:**» sigue mayúscula."""
        mayus = False
        for trozo in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", texto):
            if not trozo:
                continue
            if trozo.startswith("**"):
                p.add_run(trozo[2:-2]).bold = True
                mayus = trozo.endswith(":**")
                continue
            if mayus and not trozo.startswith("`"):
                resto = trozo.lstrip()
                trozo = trozo[: len(trozo) - len(resto)] + resto[:1].upper() + resto[1:]
            mayus = False
            if trozo.startswith("`"):
                r = p.add_run(trozo[1:-1])
                r.font.name, r.font.size = "Consolas", Pt(11)
            else:
                p.add_run(trozo)

    # ---------- bloques ----------
    def nueva_lista(self) -> str | None:
        """Numeración que reinicia en 1 sobre la definición del estilo de pasos (None si el estilo no numera)."""
        num_estilo = self.estilo("pasos").element.find(f"{qn('w:pPr')}/{qn('w:numPr')}/{qn('w:numId')}")
        if num_estilo is None:
            return None
        numbering = self.doc.part.numbering_part.element
        num = numbering.find(f"{qn('w:num')}[@{qn('w:numId')}='{num_estilo.get(qn('w:val'))}']")
        nuevo = OxmlElement("w:num")
        nuevo.set(qn("w:numId"), str(max(int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))) + 1))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), num.find(qn("w:abstractNumId")).get(qn("w:val")))
        o, so = OxmlElement("w:lvlOverride"), OxmlElement("w:startOverride")
        o.set(qn("w:ilvl"), "0")
        so.set(qn("w:val"), "1")
        o.append(so)
        nuevo.extend([a, o])
        numbering.append(nuevo)
        return nuevo.get(qn("w:numId"))

    def paso(self, num_id: str | None, texto: str) -> None:
        p = self.parrafo("pasos")
        if num_id:
            numpr = OxmlElement("w:numPr")
            for tag, val in (("w:ilvl", "0"), ("w:numId", num_id)):
                e = OxmlElement(tag)
                e.set(qn("w:val"), val)
                numpr.append(e)
            p._p.get_or_add_pPr().append(numpr)
        self.enriquecido(p, texto)

    def tabla(self, nombre: str) -> None:
        t_cfg = self.tablas[nombre]
        self.parrafo("tabla_titulo").add_run(self.titulo_tg(t_cfg["titulo"]))
        cab, filas = t_cfg["encabezados"], t_cfg["filas"]
        t = self.doc.add_table(rows=1 + len(filas), cols=len(cab))
        t.style = self.estilo("tabla")
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, h in enumerate(cab):
            t.cell(0, j).paragraphs[0].add_run(str(h)).bold = True
        for i, fila in enumerate(filas, start=1):
            for j, v in enumerate(fila):
                t.cell(i, j).paragraphs[0].add_run(str(v))
        self.nota(t_cfg["nota"])

    def glosario_(self) -> None:
        for termino, definicion in sorted(self.glosario, key=lambda f: unicodedata.normalize("NFD", f[0]).casefold()):
            self.parrafo("glosario_termino").add_run(termino)
            self.parrafo("glosario_def").add_run(definicion)

    def marcar(self, nombre: str) -> Path:
        """Copia la captura a figuras/ con sus marcadores rojos numerados y su recorte."""
        self.figuras.mkdir(parents=True, exist_ok=True)
        im = Image.open(self.capturas / f"{nombre}.png").convert("RGB")
        d = ImageDraw.Draw(im)
        try:
            f = ImageFont.truetype("DejaVuSans-Bold.ttf", 18)
        except OSError:
            f = ImageFont.load_default()
        for n, (x, y) in enumerate(self.marcadores.get(nombre, []), start=1):
            d.ellipse((x - 15, y - 15, x + 15, y + 15), fill=(215, 25, 32), outline="white", width=2)
            d.text((x, y), str(n), fill="white", font=f, anchor="mm")
        if nombre in self.recortes:
            im = im.crop(tuple(self.recortes[nombre]))
        destino = self.figuras / f"{nombre}.png"
        im.save(destino)
        return destino

    def figura(self, nombre: str, titulo: str, texto_nota: str) -> None:
        cap = self.parrafo("figura_titulo")
        cap.paragraph_format.keep_with_next = True
        cap.add_run(self.titulo_tg(titulo))
        p = self.parrafo(centro=True)
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(self.marcar(nombre)), width=Cm(self.ancho))
        self.nota(texto_nota)

    def recuadro(self, etiqueta: str, texto: str) -> None:
        """«Consejo:» y «¡Atención!» como texto normal, sin recuadro ni sombreado (como el TG)."""
        p = self.parrafo()
        p.add_run(etiqueta + " ").bold = True
        self.enriquecido(p, texto)

    def codigo(self, texto: str) -> None:
        p = self.parrafo("codigo")
        lineas = texto.strip().splitlines()
        for i, linea in enumerate(lineas):
            r = p.add_run(linea)
            if i < len(lineas) - 1:
                r.add_break()

    # ---------- estructura ----------
    def resolver_refs(self, md: str) -> str:
        figs = {f: n for n, f in enumerate(re.findall(r"\{\{FIG:\s*(\S+)", md), start=1)}
        tabs = {t: n for n, t in enumerate(re.findall(r"^\{\{TABLA-(?!GLOSARIO)([\w-]+)\}\}", md, re.MULTILINE), start=1)}
        faltan = [f"{f}.png" for f in figs if not (self.capturas / f"{f}.png").exists()]
        faltan += [f"tabla {t} (falta en tablas:)" for t in tabs if t not in self.tablas]
        if faltan:
            sys.exit(f"Faltan: {', '.join(faltan)}")

        def ref(m, n_de, etiqueta):
            if m[1] not in n_de:
                sys.exit(f"Referencia a {etiqueta.lower()} inexistente: {m[0]}")
            return f"{etiqueta} {n_de[m[1]]}"
        md = re.sub(r"\{\{REF:\s*(\S+?)\s*\}\}", lambda m: ref(m, figs, "Figura"), md)
        md = re.sub(r"\{\{REF-TABLA:\s*([\w-]+)\s*\}\}", lambda m: ref(m, tabs, "Tabla"), md)
        if "{{REF" in md:
            sys.exit("Quedó una referencia {{REF sin resolver")
        return md

    @staticmethod
    def sin_titulos_seguidos(md: str) -> None:
        """Entre un título y el siguiente va al menos un párrafo (regla de forma del TG)."""
        lineas = [x for x in md.splitlines() if x.strip()]
        for a, b in itertools.pairwise(lineas):
            if a.startswith(("## ", "### ", "#### ", "##### ")) and b.startswith("#"):
                sys.exit(f"Títulos seguidos sin párrafo entre ellos: «{a}» → «{b}»")

    @staticmethod
    def capitulos(md: str) -> dict[str, str]:
        caps, actual = {}, None
        for linea in md.splitlines():
            if linea.startswith("## "):
                m = re.match(r"CAPÍTULO \d+:\s*(.+)", linea[3:])
                actual = (m[1] if m else linea[3:]).strip()
                caps[actual] = []
            elif actual:
                caps[actual].append(linea)
        return {t: "\n".join(ls) for t, ls in caps.items()}

    def vaciar(self, h1):
        """Borra lo que sigue al H1 de la plantilla hasta el salto de sección; devuelve el ancla donde insertar."""
        body, quitar, ancla = self.doc.element.body, [], None
        el = h1.getnext()
        while el is not None and el.tag != qn("w:sectPr"):
            quitar.append(el)
            sect = el.find(f"{qn('w:pPr')}/{qn('w:sectPr')}") if el.tag == qn("w:p") else None
            if sect is not None:
                ancla = OxmlElement("w:p")
                ppr = OxmlElement("w:pPr")
                ppr.append(sect)
                ancla.append(ppr)
                el.addnext(ancla)
                break
            el = el.getnext()
        for e in quitar:
            body.remove(e)
        return ancla

    def cuerpo(self, md: str) -> None:
        lineas, lista, i = md.splitlines(), None, 0
        while i < len(lineas):
            linea = lineas[i].rstrip()
            while linea.startswith("{{") and "}}" not in linea:  # marca de varias líneas
                i += 1
                linea += "\n" + lineas[i].rstrip()
            i += 1
            if not linea or linea.startswith("# "):
                continue
            if m := re.match(r"\d+\. (.+)", linea):
                lista = lista or self.nueva_lista() or "sin-numeracion"
                self.paso(None if lista == "sin-numeracion" else lista, m[1])
                continue
            lista = None
            if m := re.fullmatch(r"\{\{FIG:\s*(\S+)\s*\|\s*(.+?)\s*\|\s*(.+?)\}\}", linea, re.DOTALL):
                self.figura(m[1], m[2].strip(), m[3].strip())
            elif m := re.fullmatch(r"\{\{CONSEJO:\s*(.+)\}\}", linea, re.DOTALL):
                self.recuadro("Consejo:", m[1].strip())
            elif m := re.fullmatch(r"\{\{ATENCION:\s*(.+)\}\}", linea, re.DOTALL):
                self.recuadro("¡Atención!", m[1].strip())
            elif m := re.fullmatch(r"\{\{CODIGO:\s*(.+)\}\}", linea, re.DOTALL):
                self.codigo(m[1])
            elif linea == "{{TABLA-GLOSARIO}}":
                self.glosario_()
            elif m := re.fullmatch(r"\{\{TABLA-([\w-]+)\}\}", linea):
                self.tabla(m[1])
            elif m := re.match(r"(#{3,5}) (.+)", linea):
                texto = re.sub(r"^\d+(\.\d+)*\.\s*", "", m[2])
                nivel = len(m[1]) - 1
                self.parrafo(f"h{nivel}").add_run(texto.upper() if nivel == 2 else self.titulo_tg(texto))
            elif linea.startswith("- "):
                self.enriquecido(self.parrafo("vinieta"), linea[2:])
            else:
                self.enriquecido(self.parrafo(), linea)

    def generar(self) -> Path:
        md = self.fuente.read_text(encoding="utf-8")
        self.sin_titulos_seguidos(md)
        caps = self.capitulos(self.resolver_refs(md))
        body = self.doc.element.body
        h1s = {p.text.strip(): p._p for p in self.doc.paragraphs if p.style.name == self.estilo("h1").name}
        if self.plantilla:
            faltan = [t for t in caps if t not in h1s]
            if faltan:
                sys.exit(f"La plantilla no tiene el capítulo {faltan} (H1 de la plantilla: {list(h1s)})")
        for titulo, texto in caps.items():
            if self.plantilla:
                ancla, n = self.vaciar(h1s[titulo]), len(body)
                self.cuerpo(texto)
                if ancla is not None:  # lo nuevo quedó al final: llevarlo a su capítulo
                    for e in list(body)[n - 1: len(body) - 1]:
                        ancla.addprevious(e)
            else:
                self.parrafo("h1").add_run(titulo)
                self.cuerpo(texto)
        usados = set(re.findall(r'r:(?:embed|id|link)="([^"]+)"', self.doc.element.xml))
        for rid, rel in list(self.doc.part.rels.items()):  # capturas viejas de la plantilla, ya sin uso
            if rel.reltype.endswith("/image") and rid not in usados:
                self.doc.part.drop_rel(rid)
        if self.doc.settings.element.find(qn("w:updateFields")) is None:  # Word ofrece actualizar los índices
            upd = OxmlElement("w:updateFields")
            upd.set(qn("w:val"), "true")
            self.doc.settings.element.append(upd)
        self.doc.save(self.salida)
        return self.salida


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    cfg_path = Path(sys.argv[1]).resolve()
    cfg = yaml.safe_load(cfg_path.read_text(encoding="utf-8"))
    print("ok", Manual(cfg, cfg_path.parent).generar())


if __name__ == "__main__":
    main()
