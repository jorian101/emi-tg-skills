"""Checks de diapositivas.py, triptico.py y tv.py con un mazo y un yaml mínimos (uv run --with python-pptx --with python-docx --with pillow --with pyyaml python este_archivo.py)."""

import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from pptx import Presentation
from pptx.util import Inches

D = Path(__file__).resolve().parent.parent


def correr(script: str, *args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [sys.executable, str(D / script), *args],
        capture_output=True,
        text=True,
        check=False,
    )


def main() -> None:
    with tempfile.TemporaryDirectory() as t:
        t = Path(t)
        prs = Presentation()
        for n in (1, 2):
            s = prs.slides.add_slide(prs.slide_layouts[6])
            s.shapes.add_textbox(
                Inches(1), Inches(1), Inches(4), Inches(1)
            ).text_frame.text = f"Resultado {n}: 26,5 % y 99,9 %"
        prs.save(t / "m.pptx")
        tg = t / "tg.md"
        tg.write_text("El tiempo bajó 26,5 % en el piloto.", encoding="utf-8")
        assert (
            correr(
                "diapositivas.py", "cifras", str(t / "m.pptx"), "--tg", str(tg)
            ).returncode
            == 1
        )  # 99,9 no está
        tg.write_text("bajó 26,5 % y 99,9 % mejoró", encoding="utf-8")
        assert (
            correr(
                "diapositivas.py", "cifras", str(t / "m.pptx"), "--tg", str(tg)
            ).returncode
            == 0
        )
        assert (
            correr("diapositivas.py", "ocultar", str(t / "m.pptx"), "2").returncode == 0
        )
        assert Presentation(t / "m.pptx").slides[1]._element.get("show") == "0"
        (t / "g.yaml").write_text(
            yaml.safe_dump([{"n": 1, "texto": "decir esto"}]), encoding="utf-8"
        )
        assert (
            correr(
                "diapositivas.py", "notas", str(t / "m.pptx"), str(t / "g.yaml")
            ).returncode
            == 0
        )
        assert (
            Presentation(t / "m.pptx").slides[0].notes_slide.notes_text_frame.text
            == "decir esto"
        )

        (t / "tv.yaml").write_text(
            yaml.safe_dump(
                {
                    "titulo": "TV",
                    "items": [{"tecla": "1", "titulo": "A", "archivo": "no.mp4"}],
                }
            ),
            encoding="utf-8",
        )
        r = correr("tv.py", str(t / "tv.yaml"))
        assert (
            r.returncode == 0
            and "por grabar" in r.stdout
            and (t / "index.html").exists()
        )

        ejemplo = D.parent.parent / "assets/defensa/triptico.example.yaml"
        (t / "img").mkdir()
        from PIL import Image

        Image.new("RGB", (40, 40), "blue").save(t / "img/logo.png")
        Image.new("RGB", (40, 40), "red").save(t / "img/foto.jpg")
        (t / "triptico.yaml").write_text(
            ejemplo.read_text(encoding="utf-8"), encoding="utf-8"
        )
        assert (
            correr("triptico.py", "generar", str(t / "triptico.yaml")).returncode == 0
        )
        assert (t / "triptico.docx").exists() and (t / "biptico.docx").exists()
        tg.write_text("nada de 00", encoding="utf-8")
        assert (
            correr(
                "triptico.py", "verificar", str(t / "triptico.yaml"), "--tg", str(tg)
            ).returncode
            == 0
        )
        # crear: el mazo sale del yaml, la oculta queda al final y el tope de palabras se hace cumplir
        from PIL import Image

        Image.new("RGB", (2400, 1350), "white").save(t / "fig.png")
        largo = " ".join(["palabra"] * 120)
        spec = {"salida": "nuevo.pptx", "diapositivas": [
            {"tipo": "portada", "titulo": "T", "caso": "C", "estudiante": "E"},
            {"tipo": "figura", "imagen": "fig.png", "leyenda": "Figura 1 · F", "anexo": "A", "oculta": True},
            {"tipo": "texto", "titulo": "OBJETIVO GENERAL", "texto": "Literal."},
        ]}
        (t / "d.yaml").write_text(yaml.safe_dump(spec, allow_unicode=True), encoding="utf-8")
        assert correr("diapositivas.py", "crear", str(t / "d.yaml")).returncode == 0
        nuevo = Presentation(t / "nuevo.pptx")
        assert [s._element.get("show") for s in nuevo.slides][-1] == "0"
        assert "Ver Anexo A" in "".join(f.text_frame.text for f in nuevo.slides[-1].shapes if f.has_text_frame)
        spec["diapositivas"].append({"tipo": "texto", "titulo": "X", "texto": largo})
        (t / "d.yaml").write_text(yaml.safe_dump(spec, allow_unicode=True), encoding="utf-8")
        assert correr("diapositivas.py", "crear", str(t / "d.yaml")).returncode == 1
        assert correr("diapositivas.py", "verificar", str(t / "nuevo.pptx")).returncode == 1
    print("ok test_diapositivas_triptico_tv")


if __name__ == "__main__":
    main()
