"""Arma el mazo de la defensa desde cero a partir de un YAML (lo llama `diapositivas.py crear SPEC.yaml`).

El YAML dice QUÉ va (el contenido sale del TG, con su anexo y sus páginas); este módulo solo pone cada cosa
en su lugar con las reglas de legibilidad de references/defensa/diapositivas.md. Modelo comentado:
assets/defensa/diapositivas.example.yaml. Tipos de diapositiva:

  portada   titulo, caso, estudiante, tutor, lugar_fecha, logos: [img]
  texto     titulo, texto | vinietas: [..], imagen (a la derecha), paginas («pp. 120-121»), anexo
  figura    imagen, leyenda, anexo                 figura a pantalla completa con una línea de leyenda
  tabla     titulo, encabezados: [..], filas: [[..]], anexo
  imagenes  titulo, imagenes: [..], pie            fotos de campo, documentación administrativa
  tarjetas  titulo, tarjetas: [{cifra, texto}], pie
  seccion   titulo, imagen                         demostración, gracias

Toda diapositiva acepta `oculta: true` (va al final, como respaldo) y `notas:` (guion del orador).
"""

from __future__ import annotations

from pathlib import Path

import yaml
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Cm, Emu, Pt

TOPE_PALABRAS = 90  # calibrado con un mazo real de 30 min: mediana 46, objetivos 86 (references/defensa/diapositivas.md)
LETRA_MIN = 14
AZUL = RGBColor.from_string("032154")
GRIS = RGBColor.from_string("495057")


def _layout(prs: Presentation, nombres: tuple[str, ...], indice: int):
    for lay in prs.slide_layouts:
        if lay.name.lower() in nombres:
            return lay
    return prs.slide_layouts[min(indice, len(prs.slide_layouts) - 1)]


def _parrafos(caja, lineas: list[str], tam: int, color=AZUL, negrita=False, centro=False, vinieta=False) -> None:
    tf = caja.text_frame
    tf.word_wrap = True
    for i, t in enumerate(lineas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = ("• " if vinieta else "") + str(t)
        p.alignment = PP_ALIGN.CENTER if centro else PP_ALIGN.LEFT
        p.space_after = Pt(6)
        for r in p.runs:
            r.font.size, r.font.bold, r.font.color.rgb = Pt(max(tam, LETRA_MIN)), negrita, color


def _texto(s, x, y, w, h, lineas, tam, **kw):
    caja = s.shapes.add_textbox(x, y, w, h)
    _parrafos(caja, [lineas] if isinstance(lineas, str) else lineas, tam, **kw)
    return caja


def _imagen(s, ruta: Path, x, y, w, h, descr: str):
    """La imagen entra entera en la caja, centrada y sin deformarse; lleva texto alternativo."""
    with Image.open(ruta) as im:
        iw, ih = im.size
    esc = min(w / iw, h / ih)
    pw, ph = int(iw * esc), int(ih * esc)
    pic = s.shapes.add_picture(str(ruta), x + (w - pw) // 2, y + (h - ph) // 2, pw, ph)
    pic._element.xpath(".//p:cNvPr")[0].set("descr", descr or ruta.stem)
    return pic


def _titulo(s, titulo: str, W) -> int:
    if s.shapes.title is not None:
        s.shapes.title.text = titulo
        return s.shapes.title.top + s.shapes.title.height
    _texto(s, Cm(1.2), Cm(0.6), W - Cm(2.4), Cm(2), titulo, 32, negrita=True)
    return Cm(2.8)


def _cita(d: dict) -> str:
    partes = [d[k] for k in ("leyenda",) if d.get(k)]
    if d.get("anexo"):
        partes.append(f"Ver Anexo {d['anexo']}")
    if d.get("paginas"):
        partes.append(d["paginas"])
    return " · ".join(partes)


def _palabras(d: dict) -> int:
    textos = [d.get("titulo", ""), d.get("texto", ""), *d.get("vinietas", []), d.get("pie", "")]
    textos += [f"{t.get('cifra', '')} {t.get('texto', '')}" for t in d.get("tarjetas", [])]
    return sum(len(str(t).split()) for t in textos)


def construir(spec: dict, base: Path) -> tuple[Presentation, list[str]]:
    plantilla = spec.get("plantilla")
    prs = Presentation(str(base / plantilla)) if plantilla else Presentation()
    if not plantilla:
        prs.slide_width, prs.slide_height = Cm(33.867), Cm(19.05)  # 16:9
    for sid in list(prs.slides._sldIdLst):  # de la plantilla solo se usan los diseños
        prs.part.drop_rel(sid.rId)
        prs.slides._sldIdLst.remove(sid)
    W, H = prs.slide_width, prs.slide_height
    tope = spec.get("tope_palabras", TOPE_PALABRAS)
    l_titulo = _layout(prs, ("diapositiva de título", "title slide"), 0)
    l_objetos = _layout(prs, ("solo el título", "title only"), 5)
    l_blanco = _layout(prs, ("en blanco", "blank"), 6)
    avisos, ocultas = [], []
    ruta = lambda r: (base / r).resolve()

    for n, d in enumerate(spec["diapositivas"], 1):
        tipo = d["tipo"]
        if (p := _palabras(d)) > tope and tipo not in ("tabla",) and not d.get("literal"):
            avisos.append(f"diapositiva {n} ({d.get('titulo', tipo)}): {p} palabras > {tope}; partirla o pasar texto a imagen")
        if tipo == "portada":
            s = prs.slides.add_slide(l_titulo)
            for ph in list(s.placeholders):
                ph._element.getparent().remove(ph._element)
            logos = d.get("logos", [])
            for i, lg in enumerate(logos):
                _imagen(s, ruta(lg), Cm(1) + i * (W - Cm(6)), Cm(0.6), Cm(4), Cm(3.2), "logo")
            y = Cm(4.2)
            for t, tam, neg in ((d["titulo"], 30, True), (f"CASO: {d['caso']}" if d.get("caso") else "", 22, True),
                                (d.get("estudiante", ""), 20, False),
                                (f"TUTOR: {d['tutor']}" if d.get("tutor") else "", 18, False),
                                (d.get("lugar_fecha", ""), 18, False)):
                if t:
                    c = _texto(s, Cm(2), y, W - Cm(4), Cm(2.6 if neg and tam == 30 else 1.4), t, tam,
                               negrita=neg, centro=True)
                    y = c.top + c.height + Cm(0.5)
        elif tipo == "figura":
            s = prs.slides.add_slide(l_blanco)
            leyenda = _cita(d)
            _imagen(s, ruta(d["imagen"]), Cm(0.6), Cm(0.4), W - Cm(1.2), H - Cm(2.2), d.get("leyenda", ""))
            _texto(s, Cm(1), H - Cm(1.7), W - Cm(2), Cm(1.2), leyenda, 14, color=GRIS, centro=True)
        else:
            s = prs.slides.add_slide(l_objetos if d.get("titulo") else l_blanco)
            y0 = _titulo(s, d["titulo"], W) + Cm(0.4) if d.get("titulo") else Cm(1)
            alto = H - y0 - Cm(1.8)
            if tipo == "texto":
                ancho = (W - Cm(2.4)) * (0.55 if d.get("imagen") else 1)
                lineas = d.get("vinietas") or [d.get("texto", "")]
                _texto(s, Cm(1.2), y0, ancho, alto, lineas, d.get("letra", 22), vinieta=bool(d.get("vinietas")))
                if d.get("imagen"):
                    _imagen(s, ruta(d["imagen"]), Cm(1.6) + ancho, y0, W - ancho - Cm(2.8), alto, d.get("titulo", ""))
            elif tipo == "tabla":
                filas = [d["encabezados"], *d["filas"]]
                letra = 16 if len(filas) <= 7 else LETRA_MIN
                tb = s.shapes.add_table(len(filas), len(filas[0]), Cm(1.2), y0, W - Cm(2.4),
                                        Emu(min(alto, Cm(1.1) * len(filas)))).table
                for i, fila in enumerate(filas):
                    for j, v in enumerate(fila):
                        c = tb.cell(i, j)
                        c.text = str(v)
                        for p in c.text_frame.paragraphs:
                            for r in p.runs:
                                r.font.size, r.font.bold = Pt(letra), i == 0
            elif tipo in ("imagenes", "seccion"):
                imgs = d.get("imagenes") or ([d["imagen"]] if d.get("imagen") else [])
                if imgs:
                    cols = min(len(imgs), 4)
                    filas_n = -(-len(imgs) // cols)
                    cw, ch = (W - Cm(2.4)) // cols, alto // filas_n
                    for i, im in enumerate(imgs):
                        _imagen(s, ruta(im), Cm(1.2) + (i % cols) * cw + Cm(0.2), y0 + (i // cols) * ch + Cm(0.2),
                                cw - Cm(0.4), ch - Cm(0.4), Path(im).stem)
            elif tipo == "tarjetas":
                ts = d["tarjetas"]
                cw = (W - Cm(2.4)) // len(ts)
                for i, t in enumerate(ts):
                    x = Cm(1.2) + i * cw
                    _texto(s, x + Cm(0.3), y0 + Cm(1), cw - Cm(0.6), Cm(3), str(t["cifra"]), 44, negrita=True, centro=True)
                    _texto(s, x + Cm(0.3), y0 + Cm(4.2), cw - Cm(0.6), Cm(4), t["texto"], 18, color=GRIS, centro=True)
            else:
                raise SystemExit(f"diapositiva {n}: tipo desconocido «{tipo}»")
            pie = d.get("pie") or _cita(d)
            if pie:
                _texto(s, Cm(1.2), H - Cm(1.6), W - Cm(2.4), Cm(1.1), pie, 13, color=GRIS)
        if d.get("notas"):
            s.notes_slide.notes_text_frame.text = d["notas"]
        if d.get("oculta"):
            ocultas.append(s)
    lista = prs.slides._sldIdLst  # las ocultas, al final: la numeración visible queda seguida
    for s in ocultas:
        s._element.set("show", "0")
        el = next(e for e in lista if prs.slides.get(int(e.get("id"))) is s)
        lista.remove(el)
        lista.append(el)
    return prs, avisos


def crear(spec_path: Path) -> int:
    spec = yaml.safe_load(spec_path.read_text(encoding="utf-8"))
    prs, avisos = construir(spec, spec_path.parent)
    salida = (spec_path.parent / spec.get("salida", "defensa.pptx")).resolve()
    prs.save(salida)
    visibles = sum(1 for s in prs.slides if s._element.get("show") != "0")
    print(f"ok {salida.name}: {len(prs.slides)} diapositivas ({visibles} visibles)")
    if visibles > 30:
        print(f"aviso {visibles} visibles para 30 minutos (uno por minuto): ocultar detalle técnico antes de recortar texto")
    print("\n".join(avisos))
    return 1 if avisos else 0
